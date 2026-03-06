from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

#for coloured formatting of the cell elements ------->>
def _parse_color(color):
    """
    Accepts:
      - None -> default black
      - RGBColor instance -> returned as-is
      - tuple/list of 3 ints (r,g,b) -> RGBColor
      - hex string '#RRGGBB' or 'RRGGBB' -> RGBColor
    """
    if color is None:
        return RGBColor(0, 0, 0)
    if isinstance(color, RGBColor):
        return color
    if isinstance(color, (tuple, list)) and len(color) == 3:
        return RGBColor(int(color[0]), int(color[1]), int(color[2]))
    if isinstance(color, str):
        s = color.lstrip('#').strip()
        if len(s) == 6:
            try:
                r = int(s[0:2], 16)
                g = int(s[2:4], 16)
                b = int(s[4:6], 16)
                return RGBColor(r, g, b)
            except Exception:
                pass
    # fallback to black
    return RGBColor(0, 0, 0)

def format_text_in_table_cell(cell, text, font_pt=10, bold=False, align=PP_ALIGN.LEFT, color="#000000"):
    """
    Insert `text` into a table cell and apply simple formatting.
    - cell: pptx table cell object
    - text: string to insert
    - font_pt: font size in points
    - bold: boolean
    - align: PP_ALIGN value
    - color: optional color for text. Accepts:
        * None (default -> black)
        * RGBColor instance
        * (r, g, b) tuple with ints 0-255
        * hex string '#RRGGBB' or 'RRGGBB'
    """
    tf = cell.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = "" if text is None else str(text)
    run.font.size = Pt(font_pt)
    run.font.bold = bold

    rgb = _parse_color(color)
    try:
        run.font.color.rgb = rgb
    except Exception:
        # best-effort: ignore color setting failures
        pass
#---------------------<<